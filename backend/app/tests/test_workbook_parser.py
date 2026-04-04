from pathlib import Path

import pytest

pytest.importorskip("openpyxl")

from openpyxl import Workbook

from app.ingestion.workbook import EXPECTED_HEADERS, parse_checklist_workbook, regulation_search_text


def test_parse_checklist_workbook_reads_rows(tmp_path: Path) -> None:
    workbook_path = tmp_path / "checklist.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Compliance Checklist"
    worksheet.append(EXPECTED_HEADERS)
    worksheet.append(
        [
            "Chemical",
            "Storage",
            "California",
            "Alameda County",
            "Berkeley",
            "City",
            "Berkeley TMD - CUPA Registration",
            "Berkeley TMD",
            "Businesses handling hazardous materials must register.",
            "Register with Berkeley TMD.",
            None,
        ]
    )
    workbook.create_sheet("Legend")
    workbook.save(workbook_path)

    rows = parse_checklist_workbook(str(workbook_path))

    assert len(rows) == 1
    assert rows[0].regulation_name == "Berkeley TMD - CUPA Registration"


def test_regulation_search_text_includes_core_fields(tmp_path: Path) -> None:
    workbook_path = tmp_path / "checklist.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Compliance Checklist"
    worksheet.append(EXPECTED_HEADERS)
    worksheet.append(
        [
            "Chemical",
            "Storage",
            "All",
            "All",
            "All",
            "Federal",
            "OSHA Hazard Communication Standard",
            "29 CFR 1910.1200",
            "Requires a written HazCom program.",
            "Maintain written HazCom program.",
            None,
        ]
    )
    workbook.save(workbook_path)

    row = parse_checklist_workbook(str(workbook_path))[0]
    search_text = regulation_search_text(row)

    assert "OSHA Hazard Communication Standard" in search_text
    assert "29 CFR 1910.1200" in search_text
    assert "Maintain written HazCom program." in search_text
