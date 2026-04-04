from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

EXPECTED_HEADERS = [
    "Industry",
    "Function",
    "State",
    "County",
    "City / Jurisdiction",
    "Level",
    "Regulation Name",
    "Regulation Code / Reference",
    "Description",
    "Action Required",
    "Status",
]


@dataclass(slots=True)
class ChecklistRow:
    industry: str
    function: str
    state: str
    county: str
    city_jurisdiction: str
    level: str
    regulation_name: str
    regulation_code_reference: str
    description: str
    action_required: str


def parse_checklist_workbook(workbook_path: str) -> list[ChecklistRow]:
    from openpyxl import load_workbook

    path = Path(workbook_path).expanduser()
    workbook = load_workbook(path, data_only=True)
    if "Compliance Checklist" not in workbook.sheetnames:
        raise ValueError("Workbook must contain a 'Compliance Checklist' sheet.")

    sheet = workbook["Compliance Checklist"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    header = list(rows[0])
    if header != EXPECTED_HEADERS:
        raise ValueError(f"Unexpected workbook header: {header}")

    parsed: list[ChecklistRow] = []
    for row in rows[1:]:
        if row is None or all(value in (None, "") for value in row):
            continue
        parsed.append(
            ChecklistRow(
                industry=str(row[0]).strip(),
                function=str(row[1]).strip(),
                state=str(row[2]).strip(),
                county=str(row[3]).strip(),
                city_jurisdiction=str(row[4]).strip(),
                level=str(row[5]).strip(),
                regulation_name=str(row[6]).strip(),
                regulation_code_reference=str(row[7]).strip(),
                description=str(row[8]).strip(),
                action_required=str(row[9]).strip(),
            )
        )
    return parsed


def regulation_search_text(row: ChecklistRow) -> str:
    return " | ".join(
        [
            f"industry: {row.industry}",
            f"function: {row.function}",
            f"state: {row.state}",
            f"county: {row.county}",
            f"city_jurisdiction: {row.city_jurisdiction}",
            f"level: {row.level}",
            f"regulation_name: {row.regulation_name}",
            f"regulation_code_reference: {row.regulation_code_reference}",
            f"description: {row.description}",
            f"action_required: {row.action_required}",
        ]
    )
